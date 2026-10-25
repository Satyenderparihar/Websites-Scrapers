import os
import sys
import json
import re
from collections import defaultdict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

PROJECT_DIR = r"D:\Antigravity - Agent\🌐 YouTube Websites Extractor Agent"
CHANNELS = ['SetupsAI', 'StudyTools', 'TrickPilot', 'PushpendraSinghAITech', 'TechBytes']

def build_master_all_channels_directory():
    print("=" * 60)
    print("🌐 Building Master All-Channels Consolidated Directory...")
    print("=" * 60)

    all_websites = {}
    channel_stats = {}
    total_scanned_videos = 0

    for ch in CHANNELS:
        json_path = os.path.join(PROJECT_DIR, f"📁 {ch}", f"💾 {ch}_websites_database.json")
        if not os.path.exists(json_path):
            json_path = os.path.join(PROJECT_DIR, f"💾 {ch}_websites_database.json")
        
        raw_path = os.path.join(PROJECT_DIR, f"💾 raw_{ch}_videos.txt")
        vid_count = 0
        if os.path.exists(raw_path):
            with open(raw_path, "r", encoding="utf-8", errors="ignore") as rf:
                vid_count = len([l for l in rf if "|" in l])
        total_scanned_videos += vid_count

        if not os.path.exists(json_path):
            continue

        with open(json_path, "r", encoding="utf-8") as jf:
            items = json.load(jf)

        channel_stats[ch] = {
            "videos": vid_count,
            "unique_websites": len(items)
        }

        for item in items:
            url = item["url"].rstrip("/").lower()
            if url not in all_websites:
                all_websites[url] = {
                    "name": item["name"],
                    "url": item["url"],
                    "category": item["category"],
                    "description": item["description"],
                    "channels": {ch: item["videos_count"]},
                    "total_videos": item["videos_count"],
                    "videos_by_channel": {ch: item["videos"]}
                }
            else:
                entry = all_websites[url]
                entry["channels"][ch] = item["videos_count"]
                entry["total_videos"] += item["videos_count"]
                entry["videos_by_channel"][ch] = item["videos"]

    print(f"📊 Aggregated {len(all_websites)} unique websites across {len(channel_stats)} channels ({total_scanned_videos} total videos).")

    # Group by category
    by_category = defaultdict(list)
    category_counts = defaultdict(int)
    sorted_websites = sorted(all_websites.values(), key=lambda x: (x["category"], -x["total_videos"], x["name"]))

    for w in sorted_websites:
        by_category[w["category"]].append(w)
        category_counts[w["category"]] += 1

    # 1. Generate Master All-Channels Markdown Report
    md_lines = [
        "# 🌐 All Channels Websites & AI Tools Master Directory\n",
        "> **All-in-One Master Consolidated Report**  ",
        f"> **Total Channels Included**: {len(channel_stats)} Channels ({', '.join(channel_stats.keys())})  ",
        f"> **Total Unique Websites (Zero Repeats)**: **{len(all_websites)} Websites** across all channels  ",
        f"> **Total YouTube Videos / Shorts Scanned**: **{total_scanned_videos} Videos**  ",
        "> **Links Standard**: 100% Exact Official Website URLs (Zero Google Search Links)  ",
        "> **Deduplication Standard**: Har website poori directory me sirf **1 baar** aayegi, sabhi channels ke use-cases ek hi jagah list hain.  \n",
        "## 📺 Channels Overview Table\n",
        "| Channel Name | Total Videos Scanned | Unique Websites | Dedicated Report | Dedicated Excel |",
        "| :--- | :---: | :---: | :--- | :--- |"
    ]

    for ch, stats in channel_stats.items():
        ch_enc = f"📁%20{ch}"
        md_lines.append(
            f"| **{ch}** | **{stats['videos']}** | **{stats['unique_websites']}** | "
            f"[📖 View {ch} Report]({ch_enc}/📝%20{ch}_Websites_Categorized_Report.md) | "
            f"[📊 Download Excel]({ch_enc}/📊%20{ch}_Websites_Master_List.xlsx) |"
        )

    md_lines.append("\n---\n")
    md_lines.append("## 📊 Master Category Breakdown\n")
    md_lines.append("| Category | Unique Websites Count | Total Video Mentions |")
    md_lines.append("| :--- | :---: | :---: |")

    for cat, count in sorted(category_counts.items(), key=lambda x: x[1], reverse=True):
        cat_vids = sum(w["total_videos"] for w in by_category[cat])
        md_lines.append(f"| {cat} | **{count} websites** | **{cat_vids} video mentions** |")

    md_lines.append("\n---\n")

    # Detailed Category Tables
    for cat in sorted(by_category.keys()):
        tools = by_category[cat]
        md_lines.append(f"### {cat} ({len(tools)} Unique Websites)\n")
        md_lines.append(
            "| Website Name | Official Website Link | Featured in Channels | Total Videos | Features & Capabilities (Is Website Se Ye Sab Kar Sakte Hain) | Sample Video References |"
        )
        md_lines.append("| :--- | :--- | :--- | :---: | :--- | :--- |")

        for t in tools:
            # Channel list formatting
            ch_list_str = "<br>".join([f"• **{c}**: {cnt} vids" for c, cnt in t["channels"].items()])
            clean_desc = t["description"].replace("|", "-")
            
            # Sample video links across channels
            sample_vids = []
            for ch_name, v_list in t["videos_by_channel"].items():
                if v_list:
                    first_title, first_url = v_list[0]
                    short_t = re.sub(r'#[a-zA-Z0-9_\-]+', '', first_title).strip()[:20]
                    sample_vids.append(f"[{ch_name}: {short_t}...]({first_url})")
            v_ref_str = "<br>".join(sample_vids[:3])

            md_lines.append(
                f"| **{t['name']}** | [{t['url']}]({t['url']}) | {ch_list_str} | **{t['total_videos']}** | {clean_desc} | {v_ref_str} |"
            )
        md_lines.append("\n")

    master_md_content = "\n".join(md_lines)
    master_md_path = os.path.join(PROJECT_DIR, "📝 All_Channels_Websites_Master_Categorized_Report.md")
    with open(master_md_path, "w", encoding="utf-8") as f:
        f.write(master_md_content)
    print(f"💾 Master All-Channels Markdown saved: {master_md_path}")

    # 2. Generate Master All-Channels Excel Workbook
    print("📝 Generating Master All-Channels Excel Workbook...")
    wb = openpyxl.Workbook()
    
    # Sheet 1: Master Consolidated
    ws_master = wb.active
    ws_master.title = "All Channels Consolidated"

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    link_font = Font(name="Segoe UI", size=10, color="2563EB", underline="single")
    alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    border_side = Side(border_style="thin", color="E2E8F0")
    thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

    headers = [
        "Category",
        "Website Name",
        "Exact Official Website Link",
        "Featured Channels",
        "Total Videos Count",
        "Consolidated Features (Is Website Se Ye Sab Kar Sakte Hain)",
        "Primary Video Reference Link"
    ]
    ws_master.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws_master.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws_master.row_dimensions[1].height = 28

    for row_idx, item in enumerate(sorted_websites, 2):
        ch_str = ", ".join([f"{c} ({cnt})" for c, cnt in item["channels"].items()])
        first_vid = ""
        for ch_name, v_list in item["videos_by_channel"].items():
            if v_list:
                first_vid = v_list[0][1]
                break

        row_data = [
            item["category"],
            item["name"],
            item["url"],
            ch_str,
            item["total_videos"],
            item["description"],
            first_vid
        ]
        ws_master.append(row_data)

        fill = alt_fill if row_idx % 2 == 0 else PatternFill(fill_type=None)
        for col_num in range(1, len(headers) + 1):
            cell = ws_master.cell(row=row_idx, column=col_num)
            cell.font = data_font
            cell.border = thin_border
            if fill.fill_type:
                cell.fill = fill
            cell.alignment = Alignment(vertical="center")

            if col_num in [3, 7] and cell.value and cell.value.startswith("http"):
                cell.font = link_font
                cell.hyperlink = cell.value

        ws_master.row_dimensions[row_idx].height = 24

    col_widths = [26, 26, 36, 32, 18, 65, 40]
    for i, w in enumerate(col_widths, 1):
        ws_master.column_dimensions[get_column_letter(i)].width = w

    # Sheet 2: Channels Summary
    ws_summary = wb.create_sheet(title="Channels Summary")
    ws_summary.append(["Channel Name", "Total Videos Scanned", "Unique Websites Identified"])
    for col_num in range(1, 4):
        c = ws_summary.cell(row=1, column=col_num)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
    ws_summary.row_dimensions[1].height = 26

    for r_idx, (ch, stats) in enumerate(channel_stats.items(), 2):
        ws_summary.append([ch, stats["videos"], stats["unique_websites"]])
        for col_num in range(1, 4):
            c = ws_summary.cell(row=r_idx, column=col_num)
            c.font = data_font
            c.border = thin_border
            c.alignment = Alignment(vertical="center", horizontal="center" if col_num > 1 else "left")
        ws_summary.row_dimensions[r_idx].height = 22

    ws_summary.column_dimensions['A'].width = 30
    ws_summary.column_dimensions['B'].width = 24
    ws_summary.column_dimensions['C'].width = 28

    master_excel_path = os.path.join(PROJECT_DIR, "📊 All_Channels_Websites_Master_List.xlsx")
    wb.save(master_excel_path)
    print(f"💾 Master All-Channels Excel saved: {master_excel_path}")

    # 3. Generate Master All-Channels JSON Database
    master_json_path = os.path.join(PROJECT_DIR, "💾 all_channels_websites_database.json")
    with open(master_json_path, "w", encoding="utf-8") as jf:
        json.dump(sorted_websites, jf, ensure_ascii=False, indent=2)
    print(f"💾 Master All-Channels JSON saved: {master_json_path}")

    return len(all_websites), total_scanned_videos

if __name__ == "__main__":
    build_master_all_channels_directory()
